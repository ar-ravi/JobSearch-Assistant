"""Excel-backed persistence and job tracker operations."""

from __future__ import annotations

import json
import csv
import io
import os
import threading
import uuid
from datetime import date, datetime, timedelta
from functools import wraps
from pathlib import Path
from typing import Any, Callable, TypeVar

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

DATA_DIR = Path(__file__).parent / "data"
DATA_FILE = DATA_DIR / "job_tracker.xlsx"
LEGACY_DATA_FILE = DATA_DIR / "job_tracker.json"
TEMP_FILE = DATA_DIR / "job_tracker.tmp.xlsx"
DEFAULTS: dict[str, Any] = {
    "settings": {"stale_after_days": 14},
    "companies": [],
    "applications": [],
    "activities": [],
}
STATUSES = {"watching", "saved", "applied", "screening", "interview", "offer", "rejected", "withdrawn"}
COMPANY_COLUMNS = ("id", "name", "roles_to_watch", "website", "added_date")
APPLICATION_COLUMNS = (
    "id", "company_id", "company", "role", "url", "location", "source", "status",
    "applied_date", "last_activity_date", "follow_up_date", "next_action", "notes", "created_at", "updated_at",
)
ACTIVITY_COLUMNS = ("id", "application_id", "occurred_at", "type", "status", "note")
ACTIVITY_TYPES = {"role_added", "status_change", "note", "recruiter_contact", "interview", "follow_up"}
_LOCK = threading.RLock()
F = TypeVar("F", bound=Callable[..., Any])


def _synchronized(function: F) -> F:
    @wraps(function)
    def wrapper(*args: Any, **kwargs: Any) -> Any:
        with _LOCK:
            return function(*args, **kwargs)
    return wrapper  # type: ignore[return-value]


def _today() -> str:
    return date.today().isoformat()


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value)


def _read_sheet(sheet: Any, columns: tuple[str, ...]) -> list[dict[str, Any]]:
    if sheet.max_row < 1:
        return []
    headers = [_text(cell.value).strip() for cell in sheet[1]]
    positions = {header: index for index, header in enumerate(headers) if header}
    missing = [column for column in columns if column not in positions]
    if missing:
        raise ValueError(f"The {sheet.title} sheet is missing columns: {', '.join(missing)}")
    rows: list[dict[str, Any]] = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        row = {column: values[positions[column]] if positions[column] < len(values) else None for column in columns}
        if any(value is not None for value in row.values()):
            rows.append(row)
    return rows


def _read_workbook(path: Path = DATA_FILE) -> dict[str, Any]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        required = {"Applications", "Companies", "Settings"}
        missing = required.difference(workbook.sheetnames)
        if missing:
            raise ValueError(f"Tracker workbook is missing sheets: {', '.join(sorted(missing))}")
        companies = _read_sheet(workbook["Companies"], COMPANY_COLUMNS)
        applications = _read_sheet(workbook["Applications"], APPLICATION_COLUMNS)
        activities = _read_sheet(workbook["Activity"], ACTIVITY_COLUMNS) if "Activity" in workbook.sheetnames else []
        settings_sheet = workbook["Settings"]
        settings: dict[str, Any] = {}
        for key, value in settings_sheet.iter_rows(min_row=2, max_col=2, values_only=True):
            if key:
                settings[str(key)] = value
        workbook.close()
    except (OSError, ValueError, KeyError) as exc:
        if isinstance(exc, ValueError) and str(exc).startswith("Tracker workbook"):
            raise
        raise ValueError(f"Could not read {path.name}: {exc}") from exc

    for company in companies:
        company["roles_to_watch"] = [role.strip() for role in _text(company.get("roles_to_watch")).split(";") if role.strip()]
        for key in ("id", "name", "website", "added_date"):
            company[key] = _text(company.get(key))
    for application in applications:
        for key in APPLICATION_COLUMNS:
            application[key] = _text(application.get(key))
    for activity in activities:
        for key in ACTIVITY_COLUMNS:
            activity[key] = _text(activity.get(key))
    settings.setdefault("stale_after_days", DEFAULTS["settings"]["stale_after_days"])
    return {"settings": settings, "companies": companies, "applications": applications, "activities": activities}


@_synchronized
def import_workbook(contents: bytes) -> dict[str, int]:
    """Validate and replace the tracker from an exported workbook."""
    if not contents:
        raise ValueError("Choose an Excel workbook to import")
    if len(contents) > 10 * 1024 * 1024:
        raise ValueError("The workbook must be smaller than 10 MB")
    import_path = DATA_DIR / "job_tracker.import.xlsx"
    try:
        DATA_DIR.mkdir(parents=True, exist_ok=True)
        import_path.write_bytes(contents)
        data = _read_workbook(import_path)
    finally:
        import_path.unlink(missing_ok=True)

    company_ids = [company["id"] for company in data["companies"]]
    application_ids = [application["id"] for application in data["applications"]]
    activity_ids = [activity["id"] for activity in data["activities"]]
    if len(company_ids) != len(set(company_ids)) or len(application_ids) != len(set(application_ids)) or len(activity_ids) != len(set(activity_ids)):
        raise ValueError("The workbook contains duplicate record IDs")
    if any(not company["id"] or not company["name"] for company in data["companies"]):
        raise ValueError("Every company row must have an ID and company name")
    if any(not application["id"] or not application["company"] or not application["role"] for application in data["applications"]):
        raise ValueError("Every application row must have an ID, company, and role")
    if any(activity["application_id"] not in set(application_ids) for activity in data["activities"]):
        raise ValueError("An activity refers to an application that is missing from the workbook")

    save_data(data)
    return {
        "companies": len(data["companies"]),
        "applications": len(data["applications"]),
        "activities": len(data["activities"]),
    }


@_synchronized
def import_applications_csv(contents: bytes) -> dict[str, int]:
    """Add application rows from CSV without replacing the existing tracker."""
    if not contents:
        raise ValueError("Choose a CSV file to import")
    if len(contents) > 10 * 1024 * 1024:
        raise ValueError("The CSV must be smaller than 10 MB")
    try:
        text = contents.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text, newline=""))
    except (UnicodeDecodeError, csv.Error) as exc:
        raise ValueError("Could not read the CSV file. Save it as UTF-8 and try again.") from exc
    headers = {header.strip().lower() for header in (reader.fieldnames or []) if header}
    if not {"company", "role"}.issubset(headers):
        raise ValueError("CSV must include Company and Role columns")
    rows: list[dict[str, str]] = []
    for row in reader:
        normalized = {str(key or "").strip().lower(): str(value or "").strip() for key, value in row.items()}
        company = normalized.get("company", "")
        role = normalized.get("role", "")
        if not company and not role and not any(normalized.values()):
            continue
        if not company or not role:
            raise ValueError("Each CSV application row needs both Company and Role")
        status = normalized.get("status", "saved").lower() or "saved"
        if status not in STATUSES:
            raise ValueError(f"Unknown status in CSV: {status}")
        rows.append(normalized)
        if len(rows) > 5000:
            raise ValueError("CSV import is limited to 5,000 application rows")

    data = load_data()
    existing_ids = {app["id"] for app in data["applications"]}
    imported_ids: set[str] = set()
    existing_keys = {
        (app.get("company", "").casefold(), app.get("role", "").casefold(), app.get("url", "").casefold())
        for app in data["applications"]
    }
    imported = 0
    skipped = 0
    for row in rows:
        company_name, role = row["company"], row["role"]
        url = row.get("url", "")
        key = (company_name.casefold(), role.casefold(), url.casefold())
        if key in existing_keys or (row.get("id") and (row["id"] in existing_ids or row["id"] in imported_ids)):
            skipped += 1
            continue
        company = next((item for item in data["companies"] if item["name"].casefold() == company_name.casefold()), None)
        if not company:
            company = {"id": str(uuid.uuid4()), "name": company_name, "roles_to_watch": [], "website": "", "added_date": _today()}
            data["companies"].append(company)
        now = datetime.now().isoformat(timespec="seconds")
        status = row.get("status", "saved").lower() or "saved"
        application = {
            "id": row.get("id") or str(uuid.uuid4()), "company_id": company["id"], "company": company["name"], "role": role,
            "url": url, "location": row.get("location", ""), "source": row.get("source", ""), "status": status,
            "applied_date": row.get("applied_date", ""), "last_activity_date": row.get("last_activity_date", "") or _today(),
            "follow_up_date": row.get("follow_up_date", ""), "next_action": row.get("next_action", ""), "notes": row.get("notes", ""),
            "created_at": row.get("created_at", "") or now, "updated_at": now,
        }
        data["applications"].append(application)
        imported_ids.add(application["id"])
        existing_keys.add(key)
        _record_activity(data, application, "role_added", "Imported from CSV.")
        imported += 1
    if imported:
        save_data(data)
    return {"imported": imported, "skipped": skipped}


def _style_sheet(sheet: Any, columns: tuple[str, ...]) -> None:
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = f"A1:{sheet.cell(row=1, column=len(columns)).coordinate}1"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="315D40")
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 24
    for row_number in range(2, sheet.max_row + 1):
        sheet.row_dimensions[row_number].height = 22
    for index, column in enumerate(columns, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = min(max(len(column) + 3, 15), 34)


def _set_value(cell: Any, value: Any) -> None:
    if isinstance(value, (list, tuple)):
        value = "; ".join(_text(item) for item in value)
    if isinstance(value, str):
        cell.value = value
        # Keep user-entered text such as '=role' as text, never an Excel formula.
        cell.data_type = "s"
    else:
        cell.value = value
    cell.alignment = Alignment(vertical="top", wrap_text=True)


def _write_workbook(data: dict[str, Any]) -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    applications_sheet = workbook.active
    applications_sheet.title = "Applications"
    companies_sheet = workbook.create_sheet("Companies")
    activities_sheet = workbook.create_sheet("Activity")
    settings_sheet = workbook.create_sheet("Settings")

    for sheet, columns in ((applications_sheet, APPLICATION_COLUMNS), (companies_sheet, COMPANY_COLUMNS)):
        sheet.append(columns)
        for row in data["applications"] if sheet is applications_sheet else data["companies"]:
            values = []
            for column in columns:
                value = row.get(column)
                if column == "roles_to_watch":
                    value = "; ".join(row.get(column, []))
                values.append(value)
            sheet.append(values)
        _style_sheet(sheet, columns)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        if sheet.max_row > 1:
            sheet.auto_filter.ref = f"A1:{sheet.cell(row=sheet.max_row, column=len(columns)).coordinate}"

    activities_sheet.append(ACTIVITY_COLUMNS)
    for activity in data.get("activities", []):
        activities_sheet.append([activity.get(column) for column in ACTIVITY_COLUMNS])
    _style_sheet(activities_sheet, ACTIVITY_COLUMNS)
    for row in activities_sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column, width in {"A": 38, "B": 38, "C": 23, "D": 22, "E": 18, "F": 70}.items():
        activities_sheet.column_dimensions[column].width = width
    if activities_sheet.max_row > 1:
        activities_sheet.auto_filter.ref = f"A1:F{activities_sheet.max_row}"

    settings_sheet.append(("setting", "value"))
    for key, value in data["settings"].items():
        settings_sheet.append((key, value))
    _style_sheet(settings_sheet, ("setting", "value"))
    settings_sheet.column_dimensions["A"].width = 24
    settings_sheet.column_dimensions["B"].width = 28
    companies_sheet.column_dimensions["A"].width = 38
    companies_sheet.column_dimensions["B"].width = 30
    companies_sheet.column_dimensions["C"].width = 42
    companies_sheet.column_dimensions["D"].width = 42
    companies_sheet.column_dimensions["E"].width = 16

    try:
        workbook.save(TEMP_FILE)
        workbook.close()
        os.replace(TEMP_FILE, DATA_FILE)
    except PermissionError as exc:
        TEMP_FILE.unlink(missing_ok=True)
        raise ValueError("Could not save the Excel tracker. Close it in Excel and try again.") from exc
    return None


@_synchronized
def load_data() -> dict[str, Any]:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    if DATA_FILE.exists():
        return _read_workbook()

    # Import the previous local JSON store once, then make Excel authoritative.
    if LEGACY_DATA_FILE.exists():
        try:
            data = json.loads(LEGACY_DATA_FILE.read_text(encoding="utf-8"))
        except (json.JSONDecodeError, OSError) as exc:
            raise ValueError(f"Could not migrate {LEGACY_DATA_FILE.name}: {exc}") from exc
        for key, value in DEFAULTS.items():
            data.setdefault(key, value.copy() if isinstance(value, dict) else list(value))
    else:
        data = {"settings": dict(DEFAULTS["settings"]), "companies": [], "applications": [], "activities": []}
    data.setdefault("activities", [])
    _write_workbook(data)
    return data


@_synchronized
def save_data(data: dict[str, Any]) -> dict[str, Any]:
    _write_workbook(data)
    return data


@_synchronized
def dashboard() -> dict[str, Any]:
    data = load_data()
    apps = data["applications"]
    activity_by_application: dict[str, list[dict[str, Any]]] = {}
    for activity in data.get("activities", []):
        activity_by_application.setdefault(activity.get("application_id", ""), []).append(activity)
    for app in apps:
        app["activities"] = sorted(
            activity_by_application.get(app["id"], []),
            key=lambda activity: activity.get("occurred_at", ""),
            reverse=True,
        )
    stale = get_stale_applications()
    followups = get_follow_up_applications()
    active = [a for a in apps if a.get("status") not in {"offer", "rejected", "withdrawn"}]
    return {
        "companies": data["companies"],
        "applications": sorted(apps, key=lambda a: a.get("updated_at", ""), reverse=True),
        "stats": {
            "active": len(active),
            "interviews": sum(a.get("status") == "interview" for a in apps),
            "offers": sum(a.get("status") == "offer" for a in apps),
            "needs_follow_up": len(followups),
        },
        "stale": stale,
        "followups": followups,
        "settings": data["settings"],
    }


@_synchronized
def add_company(name: str, roles_to_watch: list[str] | None = None, website: str = "") -> dict[str, Any]:
    name = name.strip()
    if not name:
        raise ValueError("Company name is required")
    data = load_data()
    existing = next((c for c in data["companies"] if c["name"].casefold() == name.casefold()), None)
    if existing:
        return existing
    company = {"id": str(uuid.uuid4()), "name": name, "roles_to_watch": roles_to_watch or [], "website": website.strip(), "added_date": _today()}
    data["companies"].append(company)
    save_data(data)
    return company


@_synchronized
def add_application(payload: dict[str, Any]) -> dict[str, Any]:
    company_name = str(payload.get("company", "")).strip()
    role = str(payload.get("role", "")).strip()
    if not company_name or not role:
        raise ValueError("Company and role are required")
    status = str(payload.get("status", "saved")).lower()
    if status not in STATUSES:
        raise ValueError(f"Status must be one of: {', '.join(sorted(STATUSES))}")
    data = load_data()
    company = next((c for c in data["companies"] if c["name"].casefold() == company_name.casefold()), None)
    if not company:
        company = add_company(company_name)
        data = load_data()
    today = _today()
    application = {
        "id": str(uuid.uuid4()), "company_id": company["id"], "company": company["name"], "role": role,
        "url": str(payload.get("url", "")).strip(), "location": str(payload.get("location", "")).strip(),
        "source": str(payload.get("source", "")).strip(), "status": status,
        "applied_date": str(payload.get("applied_date", today if status not in {"watching", "saved"} else "")),
        "last_activity_date": str(payload.get("last_activity_date", today)),
        "follow_up_date": str(payload.get("follow_up_date", "")),
        "next_action": str(payload.get("next_action", "")).strip(), "notes": str(payload.get("notes", "")).strip(),
        "created_at": datetime.now().isoformat(timespec="seconds"), "updated_at": datetime.now().isoformat(timespec="seconds"),
    }
    data["applications"].append(application)
    _record_activity(data, application, "role_added", f"Added to pipeline with status: {status_names_for_log(status)}.")
    save_data(data)
    return application


@_synchronized
def update_application(application_id: str, changes: dict[str, Any]) -> dict[str, Any]:
    data = load_data()
    app = next((a for a in data["applications"] if a["id"] == application_id), None)
    if not app:
        raise ValueError("Application not found")
    previous_status = app.get("status", "")
    allowed = {"status", "url", "location", "source", "applied_date", "last_activity_date", "follow_up_date", "next_action", "notes", "role"}
    for key, value in changes.items():
        if key in allowed:
            if key == "status" and value not in STATUSES:
                raise ValueError(f"Unknown status: {value}")
            app[key] = value
    app["updated_at"] = datetime.now().isoformat(timespec="seconds")
    if app.get("status") != previous_status:
        app["last_activity_date"] = _today()
        _record_activity(data, app, "status_change", f"Status changed from {status_names_for_log(previous_status)} to {status_names_for_log(app['status'])}.")
    save_data(data)
    return app


def status_names_for_log(status: str) -> str:
    return str(status or "Unknown").replace("_", " ").capitalize()


def _record_activity(data: dict[str, Any], app: dict[str, Any], activity_type: str, note: str) -> dict[str, Any]:
    activity = {
        "id": str(uuid.uuid4()),
        "application_id": app["id"],
        "occurred_at": datetime.now().isoformat(timespec="seconds"),
        "type": activity_type,
        "status": app.get("status", ""),
        "note": note.strip(),
    }
    data.setdefault("activities", []).append(activity)
    return activity


@_synchronized
def log_application_activity(application_id: str, activity_type: str, note: str) -> dict[str, Any]:
    activity_type = activity_type.strip().lower()
    note = note.strip()
    if activity_type not in ACTIVITY_TYPES - {"role_added", "status_change"}:
        raise ValueError("Choose a valid activity type")
    if not note:
        raise ValueError("Add a note before logging this activity")
    data = load_data()
    app = next((item for item in data["applications"] if item["id"] == application_id), None)
    if not app:
        raise ValueError("Application not found")
    activity = _record_activity(data, app, activity_type, note)
    app["last_activity_date"] = _today()
    app["updated_at"] = datetime.now().isoformat(timespec="seconds")
    save_data(data)
    return activity


@_synchronized
def get_stale_applications() -> list[dict[str, Any]]:
    data = load_data()
    threshold = int(data["settings"].get("stale_after_days", 14))
    cutoff = date.today() - timedelta(days=threshold)
    stale = []
    for app in data["applications"]:
        if app.get("status") in {"offer", "rejected", "withdrawn"}:
            continue
        scheduled = app.get("follow_up_date", "")
        try:
            if scheduled and date.fromisoformat(scheduled) > date.today():
                continue
        except ValueError:
            pass
        reference = app.get("last_activity_date") or app.get("applied_date")
        if not reference:
            continue
        try:
            activity_date = date.fromisoformat(reference)
        except ValueError:
            continue
        days = (date.today() - activity_date).days
        if activity_date <= cutoff:
            stale.append({**app, "days_since_activity": days})
    return sorted(stale, key=lambda a: a["days_since_activity"], reverse=True)


@_synchronized
def get_follow_up_applications() -> list[dict[str, Any]]:
    """Return active applications with an explicit due date or stale activity."""
    data = load_data()
    today = date.today()
    by_id: dict[str, dict[str, Any]] = {}
    has_future_follow_up: set[str] = set()
    for app in data["applications"]:
        if app.get("status") in {"offer", "rejected", "withdrawn"}:
            continue
        follow_up_date = app.get("follow_up_date", "")
        if follow_up_date:
            try:
                due = date.fromisoformat(follow_up_date)
            except ValueError:
                due = None
            if due and due > today:
                has_future_follow_up.add(app["id"])
            if due and due <= today:
                by_id[app["id"]] = {
                    **app,
                    "reminder_type": "scheduled",
                    "days_overdue": (today - due).days,
                    "reminder_date": follow_up_date,
                }
    for app in get_stale_applications():
        if app["id"] in has_future_follow_up:
            continue
        by_id.setdefault(app["id"], {
            **app,
            "reminder_type": "stale",
            "days_overdue": 0,
            "reminder_date": app.get("last_activity_date") or app.get("applied_date", ""),
        })
    return sorted(
        by_id.values(),
        key=lambda app: (
            0 if app["reminder_type"] == "scheduled" else 1,
            app.get("reminder_date", ""),
            app.get("company", "").casefold(),
        ),
    )


@_synchronized
def set_stale_after_days(days: int) -> dict[str, Any]:
    if days < 1 or days > 365:
        raise ValueError("Stale threshold must be between 1 and 365 days")
    data = load_data()
    data["settings"]["stale_after_days"] = days
    return save_data(data)["settings"]
